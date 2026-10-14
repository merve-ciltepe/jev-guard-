import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def add_hyperlink(paragraph, url, text, color="004B91", underline=True):
    """Adds a clickable hyperlink to a paragraph in python-docx."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)

    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')

    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)

    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

    new_run.append(rPr)
    new_run_text = OxmlElement('w:t')
    new_run_text.text = text
    new_run.append(new_run_text)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink

def create_report():
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header / Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("JEV x Alibaba Page-Agent: ToS Sentinel & KVKK Reflex Engine")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub.add_run("Chrome Uzantısı, Yapay Zeka Güvenlik & Hukuki Risk Analizi Proje Teslim Raporu")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Info Box
    p_info = doc.add_paragraph()
    r1 = p_info.add_run("Öğrenci: ")
    r1.bold = True
    p_info.add_run("Merve Çiltepe\n")
    r2 = p_info.add_run("Proje Türü: ")
    r2.bold = True
    p_info.add_run("Manifest V3 Chrome Extension, In-Page GUI Agent, MCP Server & Security Sandbox\n")
    r3 = p_info.add_run("Tarih: ")
    r3.bold = True
    p_info.add_run("Ekim 2026")
    p_info.paragraph_format.space_after = Pt(12)

    # Section 1: Bağlantılar
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("🔗 1. Proje ve Teslimat Bağlantıları")
    r_h1.font.color.rgb = RGBColor(2, 132, 199)

    links = [
        ("GitHub Repository:", "https://github.com/merve-ciltepe/jev-guard-", "https://github.com/merve-ciltepe/jev-guard-"),
        ("LinkedIn Gönderisi:", "https://lnkd.in/p/dVyfSQcc", "https://lnkd.in/p/dVyfSQcc"),
        ("Demo Videosu (Canlı GitHub İzleme):", "https://github.com/merve-ciltepe/jev-guard-/blob/main/cypress/videos/demo-walkthrough.mp4", "cypress/videos/demo-walkthrough.mp4"),
        ("Animasyonlu Demo (GIF Önizleme):", "https://github.com/merve-ciltepe/jev-guard-/blob/main/cypress/videos/demo-walkthrough.gif", "cypress/videos/demo-walkthrough.gif"),
        ("Mini Eval Raporu (5 Gerçek ToS):", "https://github.com/merve-ciltepe/jev-guard-/blob/main/eval/report.md", "eval/report.md"),
    ]

    for label, url, display_text in links:
        p = doc.add_paragraph(style='List Bullet')
        r_lbl = p.add_run(f"{label} ")
        r_lbl.bold = True
        add_hyperlink(p, url, display_text)
        p.paragraph_format.space_after = Pt(4)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 2: Değerlendirme Kriterleri
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("🏆 2. Değerlendirme Kriterleri Karşılama Özeti (100 Puan)")
    r_h2.font.color.rgb = RGBColor(2, 132, 199)

    # Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Kriter & Puan"
    hdr_cells[1].text = "Geliştirilen Sistem Bileşenleri"
    hdr_cells[2].text = "Sonuç / Doğrulama"

    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '0F172A')
        tcPr.append(shd)

    data = [
        (
            "Çalışan Demo\n(30 Puan)",
            "• Alibaba page-agent entegrasyonu (canlı DOM metin okuma)\n• JEV hızlı refleks yüzen hapı ve slide-over detay çekmecesi\n• Canlı DOM üzerinde birebir alıntı vurgulayıcı (Highlighter)\n• Cypress testleri ile kaydedilmiş ≤90 sn demo videosu",
            "✅ Eksiksiz Çalışıyor\n(Mutlu yol + Dürüstlük senaryosu kaydedildi)"
        ),
        (
            "Eval Kalitesi\n(30 Puan)",
            "• 5 gerçek dünya ToS seti (Meta, Zoom AI, TikTok Biyometrik, Adobe Abonelik, Adversarial SaaS)\n• Deterministik alıntı doğrulama: 'Alıntısı olmayan iddia geçersizdir'\n• Zod katı JSON şeması zorlaması",
            "✅ TP: 8 | FN: 1 | FP: 0\n• Precision: %100\n• Recall: %88.9\n• F1 Skoru: %94.1\n• Alıntı Doğrulama: %100"
        ),
        (
            "Siber Güvenlik\n(20 Puan)",
            "• Zorunlu Prompt Injection Testi tamamlandı\n• Sözleşmeye gizlenen sistem komutu manipülasyonları\n• HTML yorum ve komut filtresi + Kriptografik Nonce Sandboxing\n• Kırmızı alarm acil durum refleksi ve skora ceza ekleme",
            "✅ %100 Saldırı Engelleme\n(Adversarial sistem manipülasyonu başarıyla nötralize edildi)"
        ),
        (
            "Mimari & Kod\n(20 Puan)",
            "• Model Context Protocol (MCP) sunucusu entegrasyonu\n• KVKK 6698 (Madde 4, 5, 6, 8-9, 11) uyum kural seti\n• Kullanıcı savunma tercihleri (kırmızı çizgiler)\n• Manifest V3 standartlarına tam uyumlu TypeScript mimarisi",
            "✅ %100 Modüler Mimari\n• BYO LLM Key desteği\n• 'Hukuki tavsiye vermez' yasal uyarısı arayüzde zorunlu"
        )
    ]

    for row_data in data:
        row_cells = table.add_row().cells
        for i, text in enumerate(row_data):
            row_cells[i].text = text
            row_cells[i].paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: Zorunlu Yasal Uyarı
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("⚠️ 3. Zorunlu Yasal Uyarı (Legal Disclaimer)")
    r_h3.font.color.rgb = RGBColor(220, 38, 38)

    p_disc = doc.add_paragraph()
    r_disc = p_disc.add_run(
        "UYGULAMA HUKUKİ TAVSİYE VERMEZ.\n"
        "Bu Chrome uzantısı, yalnızca yapay zeka destekli algoritmik risk analizi ve farkındalık aracıdır. "
        "Resmi bir hukuki danışmanlık veya avukatlık hizmeti yerine geçmez. "
        "Bu ifade uzantının Popup, Options ve web sayfası içi Jev refleks çekmecesi arayüzlerinde belirgin şekilde kullanıcıya gösterilmektedir."
    )
    r_disc.font.size = Pt(10)
    r_disc.font.italic = True
    r_disc.font.color.rgb = RGBColor(100, 116, 139)

    # Save to project folder and Desktop
    project_path = os.path.abspath("JEV_Guard_Proje_Teslim_Raporu.docx")
    desktop_path = os.path.expanduser("~/Desktop/JEV_Guard_Proje_Teslim_Raporu.docx")

    doc.save(project_path)
    try:
        doc.save(desktop_path)
        print("Desktop copy saved:", desktop_path)
    except Exception as e:
        print("Could not save to desktop:", e)

    print("Project copy saved:", project_path)

if __name__ == '__main__':
    create_report()
